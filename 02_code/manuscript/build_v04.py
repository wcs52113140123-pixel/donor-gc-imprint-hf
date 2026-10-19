# Builds Manuscript_JACC-BTS v0.4 (docx) from v04/manuscript_v0.4.md using the LetPub template styles.
# Requires in namespace: JR (dict old_ref_number -> JACC-formatted reference string), docx, re, copy, qn.
import re, copy, json, docx
from docx.oxml.ns import qn

MD = "v04/manuscript_v0.4.md"; TEMPLATE = "v03/sub/template.docx"; V03 = "v03/sub/Manuscript_JACC-BTS.docx"; OUT = "v04/Manuscript_JACC-BTS_v0.4.docx"
CITE = re.compile(r"\[@([0-9,\-\s]+)\]")

def expand(spec):
    out = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part: a, b = map(int, part.split("-")); out += list(range(a, b + 1))
        elif part: out.append(int(part))
    return out

def compress(nums):
    nums = sorted(set(nums)); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(f"{nums[i]}-{nums[j]}" if j - i >= 2 else ",".join(str(n) for n in nums[i:j + 1]))
        i = j + 1
    return ",".join(out)

def split_sections(md):
    """Return list of (h1, [lines]) blocks."""
    blocks, cur, buf = [], None, []
    for l in md.splitlines():
        if l.startswith("# "):
            if cur is not None: blocks.append((cur, buf))
            cur, buf = l[2:].strip(), []
        else: buf.append(l)
    blocks.append((cur, buf)); return blocks

def renumber(md):
    order = []
    body_start = md.index("# Introduction")
    for m in CITE.finditer(md[body_start:]):
        for n in expand(m.group(1)):
            if n not in order: order.append(n)
    mp = {old: i + 1 for i, old in enumerate(order)}
    md2 = CITE.sub(lambda m: "<sup>" + compress([mp[n] for n in expand(m.group(1))]) + "</sup>", md)
    # "reference <sup>3</sup>" in legends -> "reference 3"
    md2 = re.sub(r"reference <sup>([0-9]+)</sup>", r"reference \1", md2)
    return md2, order, mp

def add_runs(p, text):
    for tok in re.split(r"(\*\*[^*]+\*\*|<sup>[^<]+</sup>)", text):
        if not tok: continue
        if tok.startswith("**"): p.add_run(tok[2:-2]).bold = True
        elif tok.startswith("<sup>"): r = p.add_run(tok[5:-6]); r.font.superscript = True
        else: p.add_run(tok)
    return p

def no_num(p):
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr(); numPr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); numPr.append(e)
    pPr.insert(0, numPr)

def build(JR, brief_title, data_availability, tweet=None):
    md, order, mp = renumber(open(MD, encoding="utf-8").read())
    D = docx.Document(TEMPLATE); body = D.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"): body.remove(el)
    sect = body.find(qn("w:sectPr")); src = docx.Document(V03).paragraphs
    pagebreak = copy.deepcopy(src[10]._p)
    def P(text, style, numoff=False):
        p = D.add_paragraph(style=style); add_runs(p, text)
        if numoff: no_num(p)
        return p
    def PB(): sect.addprevious(copy.deepcopy(pagebreak))
    # title page copied from v0.3 (placeholders preserved)
    for i in range(10): sect.addprevious(copy.deepcopy(src[i]._p))
    PB()
    blocks = split_sections(md); words = {"text": 0, "refs": 0, "legends": 0}; in_text = False
    for h1, lines in blocks:
        if h1 is None: continue
        paras = [l.strip() for l in lines if l.strip()]
        if h1 == "Highlights":
            P(h1, "LetPub_2.1_Heading1", True); [P(x[2:], "LetPub_3.5_ListBulleted") for x in paras]; PB()
        elif h1 == "STRUCTURED ABSTRACT":
            P(h1, "LetPub_1.5_AbstractHead")
            for x in paras:
                if x.startswith("## Keywords"): P("Keywords:", "LetPub_1.7_KeywordsHead")
                elif paras.index(x) == len(paras) - 1: P(x, "LetPub_1.8_Keywords")
                else: P(x, "LetPub_1.6_AbstractText")
        elif h1 in ("Condensed Abstract", "Abbreviations and Acronyms"):
            P(h1, "LetPub_2.1_Heading1", True); [P(x, "LetPub_3.1_Text") for x in paras]
            if h1.startswith("Abbrev"): PB()
        elif h1 in ("Introduction", "Methods", "Results", "Discussion"):
            P(h1, "LetPub_2.1_Heading1", True); words["text"] += 1
            for x in paras:
                if x.startswith("## "): P(x[3:], "LetPub_2.2_Heading2", True); words["text"] += len(x[3:].split())
                else: P(x, "LetPub_3.2_TextIndent"); words["text"] += len(re.sub(r"<sup>[^<]+</sup>", "", x).split())
        elif h1 == "Perspectives":
            P(h1, "LetPub_2.1_Heading1", True)
            for x in paras:
                if x.startswith("## "): P(x[3:], "LetPub_2.2_Heading2", True)
                else: P(x, "LetPub_3.1_Text")
            for head, txt in (("Acknowledgments", "[Optional. To be supplied by the authors.]"), ("Funding Sources", "[To be supplied by the authors.]"), ("Data and Code Availability", data_availability)):
                P(head, "LetPub_6.1_Declarations"); P(txt, "LetPub_3.1_Text")
            PB(); P("References", "LetPub_7.1_References"); words["refs"] += 1
            for i, old in enumerate(order):
                ref = f"{i + 1}. {JR[old]}"; P(ref, "LetPub_7.2_ReferenceList"); words["refs"] += len(ref.split())
            PB()
        elif h1 == "Figure Titles and Legends":
            P(h1, "LetPub_2.1_Heading1", True); words["legends"] += 4
            for x in paras: P(x, "LetPub_4.2_FigureCaption"); words["legends"] += len(re.sub(r"<sup>[^<]+</sup>|\*\*", "", x).split())
        elif h1 == "Tables":
            PB(); P(h1, "LetPub_2.1_Heading1", True)
            rows = [x for x in paras if x.startswith("|") and not re.match(r"^\|[-| ]+\|$", x)]
            for x in paras:
                if not x.startswith("|") and x.startswith("**"): P(x, "LetPub_4.2_FigureCaption")
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            T = D.add_table(rows=len(cells), cols=len(cells[0]))
            try: T.style = D.styles["Table Grid"]
            except KeyError: pass
            for i, row in enumerate(cells):
                for j, c in enumerate(row):
                    cell = T.cell(i, j); cell.text = ""; run = cell.paragraphs[0].add_run(c); run.font.size = docx.shared.Pt(9); run.bold = (i == 0)
            for x in paras:
                if not x.startswith("|") and not x.startswith("**"): P(x, "LetPub_3.1_Text")
    total = sum(words.values())
    for p in D.paragraphs:
        if p.text.startswith("Brief title:") and tweet:
            q = copy.deepcopy(p._p); p._p.addnext(q)
            qp = docx.text.paragraph.Paragraph(q, p._parent)
            for r in qp.runs[1:]: r.text = ""
            qp.runs[0].text = "Tweet: "; qp.add_run(tweet); break
    for p in D.paragraphs:
        if p.text.startswith("Brief title:"):
            for r in p.runs[1:]: r.text = ""
            p.runs[0].text = "Brief title: "; p.add_run(brief_title)
        if p.text.startswith("Total word count:"):
            for r in p.runs[1:]: r.text = ""
            p.runs[0].text = "Total word count: "; p.add_run(f"{total:,} (text from Introduction to Conclusions including headings, references and figure legends)")
    D.save(OUT)
    return {"order": order, "map": mp, "words": words, "total": total, "dropped": sorted(set(JR) - set(order))}
